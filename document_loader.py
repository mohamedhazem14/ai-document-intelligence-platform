"""Reads uploaded PDF / DOCX / TXT / CSV files and splits them into chunks."""
import os

import pandas as pd
from docx import Document as DocxDocument
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from PyPDF2 import PdfReader

from config import CHUNK_OVERLAP, CHUNK_SIZE

SUPPORTED_TYPES = ["pdf", "docx", "txt", "csv"]


def _read_text_bytes(raw: bytes) -> str:
    """Decode bytes as UTF-8, falling back to Latin-1 so odd files never crash."""
    try:
        return raw.decode("utf-8")
    except UnicodeDecodeError:
        return raw.decode("latin-1")


def _load_pdf(file, name: str) -> list[Document]:
    reader = PdfReader(file)
    pages = []
    for number, page in enumerate(reader.pages, start=1):
        text = page.extract_text()
        if text and text.strip():
            pages.append(
                Document(page_content=text, metadata={"source": name, "page": number})
            )
    return pages


def _load_docx(file, name: str) -> list[Document]:
    paragraphs = [p.text for p in DocxDocument(file).paragraphs]
    text = "\n".join(paragraphs)
    if not text.strip():
        return []
    return [Document(page_content=text, metadata={"source": name, "page": 1})]


def _load_txt(file, name: str) -> list[Document]:
    text = _read_text_bytes(file.getvalue())
    if not text.strip():
        return []
    return [Document(page_content=text, metadata={"source": name, "page": 1})]


def _load_csv(file, name: str) -> list[Document]:
    text = pd.read_csv(file).to_string()
    if not text.strip():
        return []
    return [Document(page_content=text, metadata={"source": name, "page": 1})]


_LOADERS = {
    ".pdf": _load_pdf,
    ".docx": _load_docx,
    ".txt": _load_txt,
    ".csv": _load_csv,
}


def process_uploaded_files(uploaded_files) -> list[Document]:
    """Extract text from every uploaded file and return chunked LangChain Documents."""
    raw_docs: list[Document] = []

    for uploaded in uploaded_files:
        extension = os.path.splitext(uploaded.name)[1].lower()
        loader = _LOADERS.get(extension)
        if loader is None:
            continue
        raw_docs.extend(loader(uploaded, uploaded.name))

    if not raw_docs:
        raise ValueError("No readable text was found in the uploaded files.")

    splitter = RecursiveCharacterTextSplitter(
        chunk_size=CHUNK_SIZE, chunk_overlap=CHUNK_OVERLAP
    )
    chunks = splitter.split_documents(raw_docs)

    for index, chunk in enumerate(chunks, start=1):
        chunk.metadata["chunk_id"] = f"chunk_{index}"

    return chunks
